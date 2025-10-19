
from openpyxl import load_workbook

class MasterWriter:
    def __init__(self, master_path: str):
        self.master_path = master_path
        self.wb = load_workbook(master_path)
        # Choose the right sheet—Codex: bind to your real sheet name
        self.ws = self.wb.active
        # Build header map (header text -> column index)
        self.headers = {}
        for cell in next(self.ws.iter_rows(min_row=1, max_row=1)):
            self.headers[str(cell.value).strip().lower()] = cell.column

    def col(self, header):
        return self.headers.get(header.strip().lower())

    def write_records(self, records, out_path: str):
        # Codex: implement row allocation (one row per reservation/guest per your 10/19 convention)
        row = 2
        for rec in records:
            self._write_row(row, rec)
            row += 1
        self.wb.save(out_path)

    def _write_row(self, row, rec):
        # Map normalized keys to real headers in MASTER.
        bindings = {
            "guest_names": "Guest Name",
            "arrival_date": "Arrival Date",
            "departure_date": "Departure Date",
            "nights": "Nights",
            "room_type": "Room Type",
            "room_number": "Room #",
            "harvest_time": "Harvest",
            "occasion": "Occasion",
            "allergies": "Allergies",
            "transportation": "Transportation",
            "activities": "Activities",
        }
        for key, header in bindings.items():
            col = self.col(header)
            if col:
                self.ws.cell(row=row, column=col, value=rec.get(key, ""))
